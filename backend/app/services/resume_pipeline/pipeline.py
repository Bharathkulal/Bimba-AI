import io
import re
from typing import Dict, Any, List, Optional

from app.core.exceptions import OCRException
from app.core.logging_service import log_stage, log_error
from app.services.resume_pipeline.models import (
    NormalizedDocument, DocumentPage, TextBlock, TableBlock, BoundingBox, SpanInfo
)
from app.services.resume_pipeline.document_detector import DocumentDetector
from app.services.resume_pipeline.layout_analyzer import LayoutAnalyzer
from app.services.resume_pipeline.table_extractor import TableExtractor
from app.services.resume_pipeline.ocr_service import PipelineOCRService
from app.services.resume_pipeline.section_detector import SectionDetector
from app.services.resume_pipeline.parsers import (
    PersonalParser, EducationParser, ExperienceParser, SkillsParser,
    ProjectParser, CertificationParser, PublicationParser, AchievementParser,
    PersonalDetailsParser, AdditionalParser
)
from app.services.resume_pipeline.confidence_scorer import ConfidenceScorer
from app.services.resume_pipeline.content_coverage import ContentCoverageValidator

try:
    import pymupdf
except ImportError:
    try:
        import fitz as pymupdf
    except ImportError:
        pymupdf = None

try:
    import docx
except ImportError:
    docx = None


def clean_text_artifacts(raw_text: str) -> str:
    """Standardize unicode characters, bullets, and dashes."""
    if not raw_text:
        return ""
    text = raw_text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[\u2010\u2011\u2012\u2013\u2014\u2015\u2212\ufffd]', '-', text)
    text = re.sub(r'[\u2018\u2019\u201b]', "'", text)
    text = re.sub(r'[\u201c\u201d\u201f]', '"', text)
    text = re.sub(r'[\u007f\u2022\u25cf\u25cb\u25a0\u25a1\uf0b7\u25ba\u2192]', '', text)
    return text


def despace_spaced_text(text: str) -> str:
    """Repairs spaced-out text e.g. 'P Y T H O N' -> 'PYTHON'."""
    lines = text.split("\n")
    cleaned_lines = []
    for line in lines:
        if re.search(r'(?:\b[A-Za-z0-9]\s){3,}', line):
            words = re.split(r'\s{2,}', line.strip())
            despaced_words = []
            for w in words:
                w_strip = w.strip()
                if re.match(r'^(?:[A-Za-z0-9]\s)+[A-Za-z0-9]$', w_strip):
                    despaced_words.append(w_strip.replace(" ", ""))
                elif len(w_strip) <= 3 and re.match(r'^(?:[A-Za-z0-9]\s?)+$', w_strip):
                    despaced_words.append(w_strip.replace(" ", ""))
                else:
                    sub_w = re.sub(r'\b([A-Za-z0-9])\s+(?=[A-Za-z0-9]\b)', r'\1', w_strip)
                    despaced_words.append(sub_w)
            cleaned_lines.append(" ".join(despaced_words))
        else:
            cleaned_lines.append(line)
    return "\n".join(cleaned_lines)


class GenericResumePipeline:
    """
    Unified Orchestrator for the Bimba AI Generic Resume Parsing Pipeline.
    Supports PDF (Text, Scanned, Mixed), DOCX, TXT with multi-column reading order,
    modular section parsers, table extraction, and 100% data preservation.
    """

    @classmethod
    def ingest_and_parse(cls, file_content: bytes, filename: str) -> Dict[str, Any]:
        log_stage("PIPELINE", "START", f"Processing resume: {filename} ({len(file_content)} bytes)")

        # 1. Document Detection & Strategy
        detection = DocumentDetector.detect_document_type(file_content, filename)
        file_type = detection.get("file_type", "pdf")

        # 2. Extract into Normalized Document Model
        if file_type == "pdf":
            norm_doc = cls._extract_pdf_to_model(file_content, filename, detection)
        elif file_type in ["docx", "doc"]:
            norm_doc = cls._extract_docx_to_model(file_content, filename)
        elif file_type == "txt":
            norm_doc = cls._extract_txt_to_model(file_content, filename)
        else:
            raise OCRException(f"Unsupported document format: {file_type}")

        # 3. Clean and Partition Lines into Sections
        raw_full_text = norm_doc.full_text
        clean_full_text = despace_spaced_text(clean_text_artifacts(raw_full_text))
        lines = [l.strip() for l in clean_full_text.split("\n") if l.strip()]

        # Filter out standalone page numbers
        filtered_lines = []
        for l in lines:
            if re.match(r'^(page\s+\d+(\s+of\s+\d+)?|\d+\s*/\s*\d+|\d+|---\s*page\s+\d+\s*---)$', l, re.IGNORECASE):
                continue
            filtered_lines.append(l)

        section_map = SectionDetector.partition_into_sections(filtered_lines)

        # 4. Modular Structured Parsing
        parsed_structured = cls._parse_sections(section_map, filtered_lines, clean_full_text)

        # 5. Calculate Confidence & Content Preservation Scores
        confidence_report = ConfidenceScorer.calculate_scores(
            clean_full_text,
            parsed_structured,
            raw_extraction_confidence=norm_doc.extraction_confidence
        )
        coverage_report = ContentCoverageValidator.validate_entities_and_coverage(clean_full_text, parsed_structured)

        parsed_structured["extraction_confidence_score"] = confidence_report["extraction_confidence_score"]
        parsed_structured["content_preservation_score"] = confidence_report["content_preservation_score"]
        parsed_structured["needs_human_review"] = confidence_report["needs_human_review"]
        parsed_structured["review_reasons"] = confidence_report["review_reasons"]

        log_stage("PIPELINE", "COMPLETED", f"Extraction Confidence: {confidence_report['extraction_confidence_score']}%, Preservation: {confidence_report['content_preservation_score']}%")

        return {
            "text": clean_full_text,
            "raw_text": clean_full_text,
            "pages": norm_doc.page_count,
            "confidence": float(round(norm_doc.extraction_confidence, 2)),
            "method": norm_doc.extraction_method,
            "structured_data": parsed_structured,
            "scores": {
                "extraction_confidence": confidence_report["extraction_confidence_score"],
                "content_preservation": confidence_report["content_preservation_score"]
            },
            "validation_report": coverage_report,
            "pages_metadata": [
                {
                    "page_number": p.page_number,
                    "confidence": p.confidence,
                    "text": p.raw_text
                }
                for p in norm_doc.pages
            ]
        }

    @classmethod
    def _extract_pdf_to_model(cls, file_content: bytes, filename: str, detection: Dict[str, Any]) -> NormalizedDocument:
        tables_by_page = TableExtractor.extract_pdf_tables(file_content)
        pages_list: List[DocumentPage] = []
        all_tables: List[TableBlock] = []
        full_text_parts: List[str] = []
        method_log = ["PyMuPDF"]

        if pymupdf is not None:
            try:
                doc = pymupdf.open(stream=file_content, filetype="pdf")
                for page_idx, page in enumerate(doc):
                    page_num = page_idx + 1
                    page_tables = tables_by_page.get(page_idx, [])
                    all_tables.extend(page_tables)

                    # Extract page dictionary with bounding boxes
                    try:
                        page_dict = page.get_text("dict")
                    except Exception:
                        page_dict = None

                    page_blocks: List[TextBlock] = []
                    if page_dict and isinstance(page_dict, dict):
                        for b_idx, b in enumerate(page_dict.get("blocks", [])):
                            b_text = ""
                            spans_info: List[SpanInfo] = []
                            max_font_size = 0.0
                            is_bold = False

                            for line in b.get("lines", []):
                                for span in line.get("spans", []):
                                    span_t = span.get("text", "").strip()
                                    if span_t:
                                        b_text += (span_t + " ")
                                        f_size = float(span.get("size", 0.0))
                                        f_name = span.get("font", "").lower()
                                        span_bold = "bold" in f_name or "black" in f_name or span.get("flags", 0) & 2 != 0
                                        if f_size > max_font_size:
                                            max_font_size = f_size
                                        if span_bold:
                                            is_bold = True
                                        spans_info.append(SpanInfo(
                                            text=span_t,
                                            font=span.get("font", ""),
                                            size=f_size,
                                            flags=span.get("flags", 0),
                                            is_bold=span_bold
                                        ))

                            b_text = b_text.strip()
                            if b_text:
                                bbox_tuple = b.get("bbox")
                                bbox_obj = None
                                if bbox_tuple and len(bbox_tuple) == 4:
                                    bbox_obj = BoundingBox(
                                        x0=float(bbox_tuple[0]),
                                        y0=float(bbox_tuple[1]),
                                        x1=float(bbox_tuple[2]),
                                        y1=float(bbox_tuple[3])
                                    )

                                # Check if block is inside extracted table
                                if bbox_obj and TableExtractor.is_bbox_inside_tables(bbox_obj, page_tables):
                                    continue

                                is_heading = max_font_size >= 11.5 or is_bold or b_text.isupper()

                                page_blocks.append(TextBlock(
                                    id=f"p{page_num}_b{b_idx+1}",
                                    page_number=page_num,
                                    text=b_text,
                                    bbox=bbox_obj,
                                    spans=spans_info,
                                    font_size=max_font_size,
                                    is_bold=is_bold,
                                    is_heading=is_heading
                                ))

                    # Layout analysis with multi-column reading order
                    page_obj = LayoutAnalyzer.analyze_page_layout(
                        page_number=page_num,
                        page_width=float(page.rect.width),
                        page_height=float(page.rect.height),
                        blocks=page_blocks,
                        tables=page_tables
                    )

                    # Check if page is scanned or text is degraded
                    p_text = page_obj.raw_text.strip()
                    char_count = len(p_text)
                    words_count = len(p_text.split())
                    alnum_count = sum(1 for c in p_text if c.isalnum())
                    conf = (alnum_count / char_count) if char_count > 0 else 0.0

                    if words_count < 15 or conf < 0.35:
                        # Trigger selective OCR
                        log_stage("OCR", "SELECTIVE_TRIGGER", f"Running OCR on low-confidence page {page_num}")
                        ocr_res = PipelineOCRService.ocr_page_pixmap(page)
                        if ocr_res.get("text"):
                            page_obj.raw_text = ocr_res["text"]
                            page_obj.is_scanned = True
                            page_obj.confidence = ocr_res.get("confidence", 0.75)
                            if "OCR" not in method_log:
                                method_log.append("OCR")

                    pages_list.append(page_obj)
                    full_text_parts.append(page_obj.raw_text)

            except Exception as e:
                log_error("PIPELINE", "PyMuPDF parsing failed, falling back to pypdf", e)

        if not pages_list:
            # Fallback to pypdf
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_content))
            p_texts = []
            for idx, p in enumerate(reader.pages):
                txt = p.extract_text() or ""
                p_texts.append(txt)
                pages_list.append(DocumentPage(
                    page_number=idx + 1,
                    width=612.0,
                    height=792.0,
                    columns_count=1,
                    blocks=[],
                    tables=[],
                    raw_text=txt
                ))
            full_text_parts = p_texts
            method_log = ["pypdf_fallback"]

        merged_text = "\n\n".join(full_text_parts).strip()
        alnum = sum(1 for c in merged_text if c.isalnum())
        total_chars = max(1, len(merged_text))
        final_conf = alnum / total_chars

        return NormalizedDocument(
            filename=filename,
            doc_type="pdf_text" if "OCR" not in method_log else "pdf_scanned",
            pages=pages_list,
            page_count=len(pages_list),
            full_text=merged_text,
            extraction_method=" + ".join(method_log),
            extraction_confidence=float(round(final_conf, 2)),
            tables=all_tables
        )

    @classmethod
    def _extract_docx_to_model(cls, file_content: bytes, filename: str) -> NormalizedDocument:
        full_text = ""
        tables: List[TableBlock] = []

        if docx is not None:
            try:
                from docx.document import Document
                from docx.oxml.table import CT_Tbl
                from docx.oxml.text.paragraph import CT_P
                from docx.table import _Cell, Table
                from docx.text.paragraph import Paragraph

                doc = docx.Document(io.BytesIO(file_content))

                def iter_block_items(parent):
                    if isinstance(parent, Document):
                        parent_elm = parent.element.body
                    elif isinstance(parent, _Cell):
                        parent_elm = parent._tc
                    else:
                        return

                    for child in parent_elm.iterchildren():
                        if isinstance(child, CT_P):
                            yield Paragraph(child, parent)
                        elif isinstance(child, CT_Tbl):
                            yield Table(child, parent)

                text_parts = []
                for block in iter_block_items(doc):
                    if isinstance(block, Paragraph):
                        p_txt = block.text.strip()
                        if p_txt:
                            is_heading = False
                            if len(p_txt) < 60:
                                if any(run.bold for run in block.runs if run.text.strip()):
                                    is_heading = True
                                if p_txt.isupper():
                                    is_heading = True
                            if is_heading:
                                text_parts.append(f"<H> {p_txt} </H>")
                            else:
                                text_parts.append(p_txt)
                    elif isinstance(block, Table):
                        t_rows = []
                        table_text = "<TABLE>\n"
                        for i, row in enumerate(block.rows):
                            row_data = ["\\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
                            if any(row_data):
                                t_rows.append(row_data)
                                if i == 0:
                                    table_text += "<TR-HEADER> " + " | ".join(row_data) + "\n"
                                else:
                                    table_text += "<TR> " + " | ".join(row_data) + "\n"
                        table_text += "</TABLE>"
                        text_parts.append(table_text)
                        if t_rows:
                            tables.append(TableBlock(
                                id=f"docx_t{len(tables)+1}",
                                page_number=1,
                                headers=t_rows[0],
                                rows=t_rows[1:] if len(t_rows) > 1 else [],
                                raw_data=t_rows
                            ))

                full_text = "\n".join(text_parts).strip()

            except Exception as e:
                log_error("PIPELINE", "DOCX extraction error, fallback to binary decode", e)
                full_text = file_content.decode('utf-8', errors='ignore')
                full_text = "".join(ch for ch in full_text if ch.isprintable() or ch in "\n\r\t ")
        else:
            full_text = file_content.decode('utf-8', errors='ignore')

        page = DocumentPage(
            page_number=1,
            width=612.0,
            height=792.0,
            columns_count=1,
            blocks=[],
            tables=tables,
            raw_text=full_text
        )

        return NormalizedDocument(
            filename=filename,
            doc_type="docx",
            pages=[page],
            page_count=1,
            full_text=full_text,
            extraction_method="python-docx",
            extraction_confidence=1.0,
            tables=tables
        )

    @classmethod
    def _extract_txt_to_model(cls, file_content: bytes, filename: str) -> NormalizedDocument:
        text = file_content.decode("utf-8", errors="ignore").strip()
        page = DocumentPage(
            page_number=1,
            width=612.0,
            height=792.0,
            columns_count=1,
            blocks=[],
            tables=[],
            raw_text=text
        )
        return NormalizedDocument(
            filename=filename,
            doc_type="txt",
            pages=[page],
            page_count=1,
            full_text=text,
            extraction_method="raw_text",
            extraction_confidence=1.0,
            tables=[]
        )

    @classmethod
    def _parse_sections(
        cls,
        section_map: Dict[str, Any],
        all_lines: List[str],
        full_text: str
    ) -> Dict[str, Any]:
        """Runs specialized modular parsers across all partitioned section lines."""
        # 1. Personal & Contact Info
        header_lines = section_map.get("header", [])
        personal_info = PersonalParser.parse(header_lines, all_lines)

        # 2. Summary & Objective
        summary_lines = section_map.get("summary", [])
        summary_text = " ".join(summary_lines).strip()

        obj_lines = section_map.get("objective", [])
        obj_text = " ".join(obj_lines).strip()

        # 3. Education
        edu_lines = section_map.get("education", [])
        education = EducationParser.parse(edu_lines)

        # 4. Experience
        exp_lines = section_map.get("experience", [])
        experience = ExperienceParser.parse(exp_lines)

        # 5. Internships
        intern_lines = section_map.get("internships", [])
        internships = ExperienceParser.parse(intern_lines)

        # 6. Technical Skills & Soft Skills
        tech_skill_lines = section_map.get("technical_skills", [])
        skills_res = SkillsParser.parse(tech_skill_lines, full_text=full_text)

        soft_skill_lines = section_map.get("soft_skills", [])
        soft_skills_res = SkillsParser.parse(soft_skill_lines)

        # 7. Projects
        proj_lines = section_map.get("projects", [])
        projects = ProjectParser.parse(proj_lines)

        # 8. Certifications
        cert_lines = section_map.get("certifications", [])
        certifications = CertificationParser.parse(cert_lines)

        # 9. Publications
        pub_lines = section_map.get("publications", [])
        publications = PublicationParser.parse(pub_lines)

        # 10. Achievements & Awards
        ach_lines = section_map.get("achievements", [])
        achievements = AchievementParser.parse(ach_lines)

        # 11. Leadership
        lead_lines = section_map.get("leadership", [])
        leadership = ExperienceParser.parse(lead_lines)

        # 12. Languages
        lang_lines = section_map.get("languages", [])
        languages = []
        for l in lang_lines:
            tokens = [t.strip() for t in re.split(r'[,|;•\t]', l) if t.strip() and len(t.strip()) < 30]
            languages.extend(tokens)

        # 13. Hobbies
        hobby_lines = section_map.get("hobbies", [])
        hobbies = []
        for h in hobby_lines:
            tokens = [t.strip() for t in re.split(r'[,|;•\t]', h) if t.strip() and len(t.strip()) < 40]
            hobbies.extend(tokens)

        # 14. Personal Details
        personal_details_lines = section_map.get("personal_details", []) + header_lines
        personal_details = PersonalDetailsParser.parse(personal_details_lines)

        # 15. Custom & Additional Sections (Zero Data Loss)
        custom_raw = section_map.get("custom_sections", [])
        additional_sections = AdditionalParser.parse(custom_raw)

        # Assemble normalized 16-section model
        return {
            "personal_info": personal_info,
            "personal_information": personal_info,
            "contact_information": personal_info,
            "summary": summary_text,
            "professional_summary": summary_text,
            "objective": obj_text,
            "career_objective": obj_text,
            "education": education,
            "work_experience": experience,
            "experience": experience,
            "internships": internships,
            "skills": skills_res["categorized"] if skills_res["categorized"] else skills_res["flat"],
            "technicalSkills": skills_res["flat"],
            "technical_skills": skills_res["flat"],
            "softSkills": soft_skills_res["flat"],
            "soft_skills": soft_skills_res["flat"],
            "personal_skills": soft_skills_res["flat"],
            "projects": projects,
            "certifications": certifications,
            "publications": publications,
            "achievements": achievements,
            "leadership_roles": leadership,
            "leadership": leadership,
            "languages": list(dict.fromkeys(languages)),
            "hobbies": list(dict.fromkeys(hobbies)),
            "hobbies_interests": list(dict.fromkeys(hobbies)),
            "personal_details": personal_details,
            "additional_sections": additional_sections,
            "custom_sections": additional_sections,
            "additional_information": additional_sections,
            "portfolioLinks": personal_info.get("other_links", []),
            "volunteerExperience": [],
            "references": []
        }
