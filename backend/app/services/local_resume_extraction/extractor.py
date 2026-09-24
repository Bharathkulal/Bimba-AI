import io
import os
import urllib.request
import logging
from typing import Dict, Any

try:
    import docx
except ImportError:
    docx = None

from app.services.pdf_extractor import PDFExtractor
from app.services.resume_extraction_service import extract_structured_data
from app.services.zero_loss_engine import ZeroLossEngine

logger = logging.getLogger("bimba_ai_pipeline")

def extract_docx_text_local(file_bytes: bytes) -> Dict[str, Any]:
    if docx is None:
        raise RuntimeError("python-docx is not installed")
    docx_file = io.BytesIO(file_bytes)
    doc = docx.Document(docx_file)
    text = ""
    for para in doc.paragraphs:
        if para.text:
            text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + " "
            text += "\n"
    return {
        "text": text.strip(),
        "pages": 1,
        "extraction_method": "python-docx"
    }

def get_file_bytes(url: str) -> bytes:
    req = urllib.request.Request(
        url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
    )
    with urllib.request.urlopen(req, timeout=15) as response:
        return response.read()

def enhance_with_spacy(text: str, parsed_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Enhance the rule-based extraction using spaCy.
    Specifically checks for PERSON entities at the top if name is missing or default.
    """
    try:
        import spacy
        nlp = spacy.load("en_core_web_sm")
        doc = nlp(text[:1000]) # check only first 1000 chars for names
        
        pi = parsed_data.get("personal_info", {})
        current_name = pi.get("name", "")
        
        if not current_name or current_name.lower() == "candidate name":
            for ent in doc.ents:
                if ent.label_ == "PERSON":
                    if len(ent.text.split()) >= 2:
                        parsed_data["personal_info"]["name"] = ent.text
                        break
    except Exception as e:
        logger.warning(f"spaCy enhancement failed or missing: {str(e)}")
        
    return parsed_data

class LocalResumeExtractor:
    @staticmethod
    def extract_from_url(url: str, filename: str) -> Dict[str, Any]:
        file_bytes = get_file_bytes(url)
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        
        if ext == "pdf":
            ext_res = PDFExtractor.extract_text_from_pdf(file_bytes, filename)
            raw_text = ext_res.get("full_text", "")
            method = ext_res.get("extraction_method", "pdf")
        elif ext in ["docx", "doc"]:
            ext_res = extract_docx_text_local(file_bytes)
            raw_text = ext_res.get("text", "")
            method = ext_res.get("extraction_method", "docx")
        else:
            raise ValueError(f"Unsupported file type: {ext}")
            
        if not raw_text.strip():
            raise ValueError("No text could be extracted from the document")
            
        # Extract structured data using deterministic rules/regex pipeline
        raw_structured = extract_structured_data(raw_text)
        
        # Enhance with spaCy where applicable
        enhanced_structured = enhance_with_spacy(raw_text, raw_structured)
        
        # Validate against Canonical ResumeData
        canonical_data = ZeroLossEngine.normalize_to_internal_model(enhanced_structured)
        
        return {
            "success": True,
            "raw_text": raw_text,
            "extracted_data": canonical_data,
            "extraction_method": method,
            "warnings": []
        }
